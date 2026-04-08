"""Bouw een xlsx-bytes-blob voor de download-knop in Streamlit."""

from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_EURO_FORMAT = '€ #,##0.00'
_UREN_FORMAT = '#,##0.00'


def build_resultaat_xlsx(df: pd.DataFrame) -> bytes:
    """Schrijf het resultaat-DataFrame naar een .xlsx in-memory en geef de bytes terug.

    Verwacht kolommen: ``Mdw.``, ``Naam``, ``Totaal uren``,
    ``Totaal aanvulling (€)``, ``Gewogen aanvulling p/u (€)``.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultaat"

    headers = list(df.columns)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="left")

    for _, row in df.iterrows():
        ws.append([row[col] for col in headers])

    # Number formats per kolom (op basis van header).
    for idx, col in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        if col == "Totaal uren":
            for cell in ws[letter][1:]:
                cell.number_format = _UREN_FORMAT
        elif "(€)" in col:
            for cell in ws[letter][1:]:
                cell.number_format = _EURO_FORMAT

    # Kolombreedtes op basis van inhoud.
    for idx, col in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        max_len = max([len(str(col))] + [len(str(v)) for v in df[col].tolist()])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 32)

    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
