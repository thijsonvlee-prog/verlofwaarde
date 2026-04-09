"""Lees verlofsaldi- en aanvulling-exports uit xlsx naar genormaliseerde DataFrames.

De headers in de salarisadministratie-exports kunnen op micro-niveau verschillen
(spaties, hoofdletters, leestekens). We normaliseren ze één keer en hernoemen
naar canonical kolomnamen die de rest van de codebase gebruikt.
"""

from __future__ import annotations

import re
from typing import IO

import pandas as pd

# Aliases per canonical kolom: alle spellingen die we accepteren in een upload.
# Sleutels worden vergeleken na normalisatie (lowercase, stripped, leestekens weg).
_VERLOFSALDI_ALIASES: dict[str, list[str]] = {
    "mdw": ["mdw", "mdw nr", "medewerker", "personeelsnr", "personeelsnummer"],
    "naam": ["naam", "medewerker naam"],
    "dv": ["dv", "dienstverband"],
    "type_verlof": ["type verlof", "verloftype", "soort verlof"],
    "bkjr": ["bkjr", "boekjaar"],
    "beginsaldo": ["beginsaldo", "begin saldo"],
}

_AANVULLING_ALIASES: dict[str, list[str]] = {
    "mdw": ["mdw", "mdw nr", "medewerker", "personeelsnr", "personeelsnummer"],
    "naam": ["naam", "medewerker naam"],
    "dv": ["dv", "dienstverband"],
    "lc": ["lc", "looncomponent"],
    "omschrijving": ["omschrijving"],
    "aanvulling_pu": ["totaalbedrag", "aanvulling per uur", "aanvulling pu"],
    "bkjr": ["bkjr", "boekjaar"],
}

_REQUIRED_VERLOFSALDI = ["mdw", "naam", "type_verlof", "beginsaldo"]
_REQUIRED_AANVULLING = ["mdw", "bkjr", "aanvulling_pu"]


def _normalize_header(name: object) -> str:
    """Lowercase, strip leestekens, normaliseer spaties — voor robuuste matching."""
    s = str(name).strip().lower()
    s = s.replace(".", "").replace(",", "")
    s = re.sub(r"\s+", " ", s)
    return s


def _build_rename_map(
    columns: list[str], aliases: dict[str, list[str]]
) -> dict[str, str]:
    """Vind voor elke canonical kolom de eerste matchende echte kolom in de upload."""
    normalized = {col: _normalize_header(col) for col in columns}
    rename: dict[str, str] = {}
    for canonical, options in aliases.items():
        wanted = {_normalize_header(o) for o in options}
        for original, norm in normalized.items():
            if norm in wanted and original not in rename:
                rename[original] = canonical
                break
    return rename


def _coerce_float(series: pd.Series) -> pd.Series:
    """Converteer naar float met tolerantie voor zowel "4.54" als "1.234,56"."""
    if pd.api.types.is_numeric_dtype(series):
        return series.astype(float)

    def _parse(value: object) -> float:
        if value is None or (isinstance(value, float) and pd.isna(value)):
            return 0.0
        if isinstance(value, (int, float)):
            return float(value)
        s = str(value).strip().replace("\u00a0", "").replace(" ", "")
        if not s:
            return 0.0
        # Nederlandse notatie: "1.234,56" → "." is duizendtal, "," is decimaal.
        # Engelse notatie: "4.54" → "." is decimaal. Onderscheid op aanwezigheid
        # van een komma.
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return 0.0

    return series.map(_parse).astype(float)


def _read_first_sheet(file: IO[bytes] | str) -> pd.DataFrame:
    try:
        return pd.read_excel(file, sheet_name=0, engine="openpyxl", dtype=object)
    except TypeError:
        # Fallback voor xlsx-bestanden met stylesheets die openpyxl niet aankan
        # (bekend probleem op Python 3.13+). read_only mode slaat stylesheet-
        # parsing over.
        if hasattr(file, "seek"):
            file.seek(0)
        from openpyxl import load_workbook

        wb = load_workbook(file, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        data = list(ws.iter_rows(values_only=True))
        wb.close()
        if not data:
            return pd.DataFrame()
        headers = [
            str(h) if h is not None else f"col_{i}"
            for i, h in enumerate(data[0])
        ]
        return pd.DataFrame(data[1:], columns=headers)


def read_verlofsaldi(file: IO[bytes] | str) -> pd.DataFrame:
    """Lees een verlofsaldi-export en geef een genormaliseerde DataFrame terug.

    Returns:
        DataFrame met kolommen ``mdw, naam, dv, type_verlof, bkjr, beginsaldo``.

    Raises:
        ValueError: als één van de verplichte kolommen ontbreekt.
    """
    raw = _read_first_sheet(file)
    rename_map = _build_rename_map(list(raw.columns), _VERLOFSALDI_ALIASES)
    df = raw.rename(columns=rename_map)

    missing = [c for c in _REQUIRED_VERLOFSALDI if c not in df.columns]
    if missing:
        raise ValueError(
            "Verlofsaldi-bestand mist verplichte kolommen: "
            + ", ".join(missing)
            + f". Gevonden kolommen: {list(raw.columns)}"
        )

    keep = [c for c in ["mdw", "naam", "dv", "type_verlof", "bkjr", "beginsaldo"] if c in df.columns]
    df = df[keep].copy()

    df["mdw"] = df["mdw"].astype(str).str.strip()
    df["naam"] = df["naam"].astype(str).str.strip()
    df["type_verlof"] = df["type_verlof"].astype(str).str.strip()
    df["beginsaldo"] = _coerce_float(df["beginsaldo"])
    if "dv" in df.columns:
        df["dv"] = df["dv"].astype(str).str.strip()
    if "bkjr" in df.columns:
        df["bkjr"] = pd.to_numeric(df["bkjr"], errors="coerce").astype("Int64")

    df = df[df["mdw"].ne("") & df["mdw"].ne("nan")]
    return df.reset_index(drop=True)


def read_aanvulling(file: IO[bytes] | str) -> pd.DataFrame:
    """Lees een aanvulling-per-uur-export en geef een genormaliseerde DataFrame terug.

    Returns:
        DataFrame met kolommen ``mdw, naam, dv, bkjr, aanvulling_pu`` (en
        eventueel ``lc``, ``omschrijving`` als die in de upload zitten).

    Raises:
        ValueError: als één van de verplichte kolommen ontbreekt.
    """
    raw = _read_first_sheet(file)
    rename_map = _build_rename_map(list(raw.columns), _AANVULLING_ALIASES)
    df = raw.rename(columns=rename_map)

    missing = [c for c in _REQUIRED_AANVULLING if c not in df.columns]
    if missing:
        raise ValueError(
            "Aanvulling-bestand mist verplichte kolommen: "
            + ", ".join(missing)
            + f". Gevonden kolommen: {list(raw.columns)}"
        )

    keep = [
        c
        for c in ["mdw", "naam", "dv", "lc", "omschrijving", "aanvulling_pu", "bkjr"]
        if c in df.columns
    ]
    df = df[keep].copy()

    df["mdw"] = df["mdw"].astype(str).str.strip()
    df["bkjr"] = pd.to_numeric(df["bkjr"], errors="coerce").astype("Int64")
    df["aanvulling_pu"] = _coerce_float(df["aanvulling_pu"])
    if "naam" in df.columns:
        df["naam"] = df["naam"].astype(str).str.strip()
    if "dv" in df.columns:
        df["dv"] = df["dv"].astype(str).str.strip()

    df = df[df["mdw"].ne("") & df["mdw"].ne("nan") & df["bkjr"].notna()]
    return df.reset_index(drop=True)
