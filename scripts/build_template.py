#!/usr/bin/env python3
"""Generate the 'Waarde vakantiedag' Excel template.

Run:
    pip install -r requirements.txt
    python scripts/build_template.py
"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUT = ROOT / "dist" / "waarde_vakantiedag_template.xlsx"

# Initial number of data rows pre-filled with formulas per input table.
# Tables auto-expand on paste; calculated columns auto-fill in Excel.
VERLOF_ROWS = 500
AANVULLING_ROWS = 500
INSTELLINGEN_ROWS = 50

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HELPER_FILL = PatternFill("solid", fgColor="FFF2CC")
HELPER_FONT = Font(bold=True, color="000000")


# ------------------------------------------------------------------ helpers

def parse_decimal(s: str) -> float:
    """Parse a Dutch-formatted decimal ('1.234,56' or '5,68') to float."""
    s = s.strip()
    if not s:
        return 0.0
    return float(s.replace(".", "").replace(",", "."))


def load_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter=";"))


def set_header(ws, row, col, text, fill=HEADER_FILL, font=HEADER_FONT):
    cell = ws.cell(row=row, column=col, value=text)
    cell.fill = fill
    cell.font = font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    return cell


def add_table(ws, name, ref):
    tbl = Table(displayName=name, name=name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tbl)


# ------------------------------------------------------------------- sheets

def write_instructie(ws):
    ws["A1"] = "Waarde vakantiedag — Gewogen aanvulling p/u"
    ws["A1"].font = Font(bold=True, size=14)

    lines = [
        "",
        "Doel",
        "    Per medewerker één gewogen gemiddelde aanvulling p/u berekenen,",
        "    zodat elk in het lopende kalenderjaar opgenomen verlofuur dezelfde",
        "    aanvulling krijgt — ongeacht uit welk opbouwjaar het verlof komt.",
        "",
        "Stappen",
        "    1. Plak de verlofsaldi-export in het blad 'Verlofsaldi'.",
        "    2. Plak de aanvulling-per-uur-export in het blad 'Aanvulling per uur'.",
        "    3. Controleer het blad 'Instellingen'. Per basis-verloftype (de",
        "       naam zonder jaartal) staat hier of er aanvulling over betaald",
        "       moet worden (Ja/Nee). Onbekende basistypes verschijnen rechts",
        "       op hetzelfde blad met status 'ONTBREEKT'.",
        "    4. Lees het resultaat af in het blad 'Resultaat'.",
        "",
        "Rekenregel",
        "    Alleen verlofregels met een 4-cijferig opbouwjaar in de naam",
        "    tellen mee. Per meetellende regel:",
        "        gewicht = Beginsaldo × aanvulling p/u (uit opbouwjaar)",
        "    Aanvullingsrechtig = Nee → gewicht = 0, maar de uren tellen wél",
        "    mee in de noemer en verlagen zo het gewogen gemiddelde.",
        "    Per medewerker (over alle DV's):",
        "        gewogen aanvulling p/u = SOM(gewicht) / SOM(Beginsaldo)",
        "",
        "Cijfervoorbeeld uit de business-specificatie",
        "    Bovenwettelijk 2025 =  32u × €3 =  €96",
        "    Wettelijk     2026 = 160u × €4 = €640",
        "    Bovenwettelijk 2026 =  60u × €4 = €240",
        "    Totaal: 252u, €976  →  gemiddeld €3,87 p/u",
        "",
        "Let op",
        "    • De aanvulling p/u wordt gezocht op Mdw. + Opbouwjaar.",
        "      De DV in het aanvullingsblad wordt genegeerd, zodat verlof uit",
        "      een eerder jaar altijd het tarief van dát jaar krijgt (ook als",
        "      dat onder een andere DV is geboekt).",
        "    • Types zonder jaartal in de naam (bv. 'Negatief verlof')",
        "      worden volledig genegeerd — niet in teller én niet in noemer.",
        "    • Basis-type met status 'ONBEKEND' in de kolom Aanvullingsrechtig",
        "      = nog niet ingesteld op het blad 'Instellingen'. Voeg die regel",
        "      dan toe en kies Ja of Nee.",
    ]
    for i, line in enumerate(lines, start=2):
        ws.cell(row=i, column=1, value=line)

    ws.column_dimensions["A"].width = 92


def write_instellingen(ws):
    set_header(ws, 1, 1, "Basis verloftype")
    set_header(ws, 1, 2, "Aanvullingsrechtig")

    defaults = [
        ("Wettelijke rechten", "Ja"),
        ("Bovenwettelijke rechten", "Ja"),
        ("Bovenwettelijke rechten (aanv.)", "Ja"),
        ("Negatief verlof", "Nee"),
    ]
    for i, (bt, ja) in enumerate(defaults, start=2):
        ws.cell(row=i, column=1, value=bt)
        ws.cell(row=i, column=2, value=ja)

    last_row = INSTELLINGEN_ROWS + 1
    add_table(ws, "tblInstellingen", f"A1:B{last_row}")

    dv = DataValidation(type="list", formula1='"Ja,Nee"', allow_blank=True)
    dv.add(f"B2:B{last_row}")
    ws.add_data_validation(dv)

    # Helper block: list all basis-types found in Verlofsaldi and show
    # whether they're already configured.
    set_header(ws, 1, 4, "Basistypes in data")
    set_header(ws, 1, 5, "Status")
    ws["D2"] = (
        '=IFERROR(SORT(UNIQUE(FILTER(tblVerlof[Basis_type],'
        'tblVerlof[Basis_type]<>""))),"")'
    )
    ws["E2"] = (
        '=IF(D2#="","",'
        'IF(ISNUMBER(MATCH(D2#,tblInstellingen[Basis verloftype],0)),'
        '"OK","ONTBREEKT"))'
    )

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["D"].width = 36
    ws.column_dimensions["E"].width = 14


def write_verlofsaldi(ws):
    plak_headers = [
        "Mdw.", "Naam", "DV", "Type verlof", "Bkjr.",
        "Actueel saldo deze periode", "Geboekt / nog op te nemen",
        "Vorige periode", "Perioderecht", "Extra perioderecht",
        "Beginsaldo", "Corr.", "Opgenomen", "Saldo",
    ]
    helper_headers = [
        "Opbouwjaar", "Basis_type", "Aanvullingsrechtig",
        "Aanvulling_pu", "Gewicht", "Meetellen",
    ]
    headers = plak_headers + helper_headers

    for col, name in enumerate(plak_headers, start=1):
        set_header(ws, 1, col, name)
    for col, name in enumerate(helper_headers, start=len(plak_headers) + 1):
        set_header(ws, 1, col, name, fill=HELPER_FILL, font=HELPER_FONT)

    fixture_rows = load_csv(ROOT / "fixtures" / "voorbeeld_verlofsaldi.csv")

    numeric_cols = {
        "Actueel saldo deze periode", "Geboekt / nog op te nemen",
        "Vorige periode", "Perioderecht", "Extra perioderecht",
        "Beginsaldo", "Corr.", "Opgenomen", "Saldo",
    }
    int_cols = {"Mdw.", "DV", "Bkjr."}

    for i, row in enumerate(fixture_rows, start=2):
        for col_idx, h in enumerate(plak_headers, start=1):
            val = row[h]
            if h in numeric_cols:
                val = parse_decimal(val)
            elif h in int_cols:
                val = int(val)
            ws.cell(row=i, column=col_idx, value=val)

    # Formulas for helper columns; written to every data row so Excel
    # recognises them as calculated columns and auto-fills on expansion.
    last_row = VERLOF_ROWS + 1
    f_opbouwjaar = (
        '=LET(t,[@[Type verlof]],'
        's,SEQUENCE(MAX(1,LEN(t)-3)),'
        'k,IFERROR(--MID(t,s,4),0),'
        'g,FILTER(k,(k>=2000)*(k<=2100),""),'
        'IFERROR(INDEX(g,1),""))'
    )
    f_basis_type = (
        '=LET(t,[@[Type verlof]],'
        'jj,[@Opbouwjaar],'
        's,IF(jj="",t,SUBSTITUTE(t," "&jj,"")),'
        'TRIM(s))'
    )
    f_aanvullingsrechtig = (
        '=IF([@[Type verlof]]="","",'
        'IFERROR(XLOOKUP([@Basis_type],'
        'tblInstellingen[Basis verloftype],'
        'tblInstellingen[Aanvullingsrechtig]),"ONBEKEND"))'
    )
    f_aanvulling_pu = (
        '=IF(OR([@Opbouwjaar]="",[@[Type verlof]]=""),0,'
        'IFERROR(INDEX(FILTER(tblAanvulling[Totaalbedrag],'
        '(tblAanvulling[Mdw.]=[@[Mdw.]])*'
        '(tblAanvulling[Bkjr.]=[@Opbouwjaar]),0),1),0))'
    )
    f_gewicht = (
        '=IF([@Aanvullingsrechtig]="Ja",'
        '[@Beginsaldo]*[@Aanvulling_pu],0)'
    )
    f_meetellen = (
        '=IF(AND([@Opbouwjaar]<>"",'
        '[@Beginsaldo]<>0,'
        'OR([@Aanvullingsrechtig]="Ja",[@Aanvullingsrechtig]="Nee")),1,0)'
    )

    for r in range(2, last_row + 1):
        ws.cell(row=r, column=15, value=f_opbouwjaar)
        ws.cell(row=r, column=16, value=f_basis_type)
        ws.cell(row=r, column=17, value=f_aanvullingsrechtig)
        ws.cell(row=r, column=18, value=f_aanvulling_pu)
        ws.cell(row=r, column=19, value=f_gewicht)
        ws.cell(row=r, column=20, value=f_meetellen)

    last_col = get_column_letter(len(headers))
    add_table(ws, "tblVerlof", f"A1:{last_col}{last_row}")

    widths = {
        "A": 8, "B": 18, "C": 6, "D": 36, "E": 8,
        "F": 12, "G": 14, "H": 12, "I": 12, "J": 14,
        "K": 12, "L": 8, "M": 12, "N": 10,
        "O": 11, "P": 30, "Q": 18, "R": 14, "S": 12, "T": 10,
    }
    for c, w in widths.items():
        ws.column_dimensions[c].width = w


def write_aanvulling(ws):
    headers = ["Mdw.", "Naam", "DV", "LC", "Omschrijving",
               "Totaalbedrag", "Bkjr."]
    for col, name in enumerate(headers, start=1):
        set_header(ws, 1, col, name)

    fixture_rows = load_csv(ROOT / "fixtures" / "voorbeeld_aanvulling.csv")

    for i, row in enumerate(fixture_rows, start=2):
        ws.cell(row=i, column=1, value=int(row["Mdw."]))
        ws.cell(row=i, column=2, value=row["Naam"])
        ws.cell(row=i, column=3, value=int(row["DV"]))
        ws.cell(row=i, column=4, value=int(row["LC"]))
        ws.cell(row=i, column=5, value=row["Omschrijving"])
        ws.cell(row=i, column=6, value=parse_decimal(row["Totaalbedrag"]))
        ws.cell(row=i, column=7, value=int(row["Bkjr."]))

    last_row = AANVULLING_ROWS + 1
    add_table(ws, "tblAanvulling", f"A1:G{last_row}")

    widths = {"A": 8, "B": 18, "C": 6, "D": 8,
              "E": 34, "F": 14, "G": 8}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w


def write_resultaat(ws):
    headers = ["Mdw.", "Naam", "Totaal uren",
               "Totaal aanvulling (€)", "Gewogen aanvulling p/u (€)"]
    for col, name in enumerate(headers, start=1):
        set_header(ws, 1, col, name)

    ws["A2"] = (
        '=IFERROR(SORT(UNIQUE(FILTER(tblVerlof[Mdw.],'
        'tblVerlof[Meetellen]=1,""))),"")'
    )
    ws["B2"] = (
        '=IF(A2#="","",'
        'IFERROR(XLOOKUP(A2#,tblVerlof[Mdw.],tblVerlof[Naam]),""))'
    )
    ws["C2"] = (
        '=IF(A2#="","",'
        'SUMIFS(tblVerlof[Beginsaldo],tblVerlof[Mdw.],A2#,'
        'tblVerlof[Meetellen],1))'
    )
    ws["D2"] = (
        '=IF(A2#="","",'
        'SUMIFS(tblVerlof[Gewicht],tblVerlof[Mdw.],A2#,'
        'tblVerlof[Meetellen],1))'
    )
    ws["E2"] = '=IF(A2#="","",IFERROR(D2#/C2#,0))'

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 28


# --------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT,
                    help="Output path for the .xlsx (default: %(default)s)")
    args = ap.parse_args()

    wb = Workbook()
    wb.remove(wb.active)

    write_instructie(wb.create_sheet("Instructie"))
    write_instellingen(wb.create_sheet("Instellingen"))
    write_verlofsaldi(wb.create_sheet("Verlofsaldi"))
    write_aanvulling(wb.create_sheet("Aanvulling per uur"))
    write_resultaat(wb.create_sheet("Resultaat"))

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
